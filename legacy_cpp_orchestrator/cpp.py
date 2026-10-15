from __future__ import annotations

import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook

from .db import client, fetch_all, chunked
from .settings import settings

SCRAPER_FILES = {
    "arogga": "arogga_basepack_scraper.py",
    "daraz": "daraz_basepack_scraper.py",
    "othoba": "othoba_basepack_scraper.py",
    "pandamart": "pandamart_basepack_scraper.py",
    "shajgoj": "shajgoj_basepack_scraper.py",
    "shwapno": "shwapno_basepack_scraper.py",
}

ACCOUNT_LABELS = {
    "arogga": "Arogga",
    "daraz": "Daraz",
    "othoba": "Othoba",
    "pandamart": "Pandamart",
    "shajgoj": "Shajgoj",
    "shwapno": "Shwapno",
}

CPP_HEADERS = [
    "Basepack", "BUSINESS UNIT", "Category", "Format", "Brand", "SKU",
    "Account Product Name", "Product URL", "Slug", "Status",
]


def _now() -> datetime:
    return datetime.now(ZoneInfo(settings.timezone))


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text


def _num(value: Any) -> float | None:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^0-9.\-]", "", str(value).replace(",", ""))
    if text in {"", ".", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _truthy_ola(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    n = _num(value)
    if n is not None:
        return n > 0
    return _clean(value).lower() in {"true", "yes", "y", "available", "in stock", "instock"}


def _normalize_daraz(value: Any) -> str:
    text = _clean(value)
    return re.sub(r"\s+", "", text.replace("_BD_", "_BD-").replace(" BD ", "_BD-"))


def _normalize_branch(value: Any) -> str:
    text = _clean(value)
    text = re.sub(r"^Delivery\s+Hero\s+Stores\s*\(Bangladesh\)\s+Limited[-\s]*", "", text, flags=re.I).strip()
    key = re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()
    aliases = {
        "bashundhara": "Bashundhara",
        "dhanmondi": "Dhanmondi",
        "gulshan": "Gulshan",
        "gulshan 2": "Gulshan 2",
        "gulshan banani": "Gulshan",
        "gulshan niketon": "Gulshan 2",
        "mirpur 02": "Mirpur 02",
        "mirpur 2": "Mirpur 02",
        "mirpur zoo road": "Mirpur 02",
        "mirpur 03": "Mirpur 03",
        "mirpur 3": "Mirpur 03",
        "mogbazar": "Mogbazar",
        "mohammadpur": "Mohammadpur",
        "rampura": "Rampura",
        "uttara": "Uttara",
        "wari": "Wari",
    }
    return aliases.get(key, text)


def _json_safe(value: Any):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _sku_for_account(code: str, value: Any) -> str:
    return _normalize_daraz(value) if code == "daraz" else _clean(value)


def _write_input_sheet(ws, code: str, rows: list[dict]) -> None:
    headers = CPP_HEADERS.copy()
    header_by_code = {
        "daraz": "daraz_sku",
        "othoba": "product_id",
        "pandamart": "Item_Id",
        "shwapno": "Code",
    }
    headers[5] = header_by_code.get(code, "SKU")
    ws.append(headers)
    for row in rows:
        basepack = row.get("basepack_name") or row.get("product_name") or f"CPP {row['account_sku']}"
        ws.append([
            basepack,
            "CPP",
            row.get("category") or "",
            "",
            row.get("brand") or "",
            row["account_sku"],
            row.get("product_name") or "",
            row.get("product_url") or "",
            "",
            "Active",
        ])


def export_cpp_inputs(targets: list[dict], output_dir: Path) -> dict[str, Path]:
    """Create scraper input files from live CPP targets, independent of OLA scope."""
    output_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for target in targets:
        grouped[target["account_code"]].append(target)

    filenames = {
        "arogga": "Arogga CPP SKU List.xlsx",
        "daraz": "Daraz CPP SKU List.xlsx",
        "othoba": "Othoba CPP SKU List.xlsx",
        "pandamart": "Pandamart CPP SKU List.xlsx",
        "shajgoj": "Shajgoj CPP SKU List.xlsx",
        "shwapno": "SHWAPNO CPP SKU List.xlsx",
    }
    result: dict[str, Path] = {}
    for code, rows in grouped.items():
        if code not in filenames or code not in SCRAPER_FILES:
            continue
        wb = Workbook()
        wb.remove(wb.active)
        if code == "pandamart":
            by_branch: dict[str, list[dict]] = defaultdict(list)
            for row in rows:
                branch = _normalize_branch(row.get("location_name") or "")
                if not branch:
                    continue
                by_branch[branch].append(row)
            for branch, branch_rows in sorted(by_branch.items()):
                ws = wb.create_sheet(branch[:31])
                _write_input_sheet(ws, code, branch_rows)
        else:
            ws = wb.create_sheet("SKU List")
            _write_input_sheet(ws, code, rows)

        path = output_dir / filenames[code]
        wb.save(path)
        result[code] = path
    return result


def _read_scrape_output(path: Path, code: str) -> list[dict]:
    xls = pd.ExcelFile(path)
    df = pd.read_excel(path, sheet_name=xls.sheet_names[0])
    rows: list[dict] = []
    for _, r in df.iterrows():
        sku = _sku_for_account(code, r.get("Account SKU"))
        if not sku:
            continue
        regular_price = _num(r.get("Price"))
        sale_price = _num(r.get("Sale Price")) if code == "shajgoj" else None
        actual_price = sale_price if sale_price is not None else regular_price
        rows.append({
            "account_sku": sku,
            "branch": _normalize_branch(r.get("Branch")) if code == "pandamart" else "",
            "product_name": _clean(r.get("Site Product Name")) or _clean(r.get("Account Product Name")),
            "actual_price": actual_price,
            "original_price": _num(r.get("Original Price")) if r.get("Original Price") is not None else regular_price,
            "in_stock": _truthy_ola(r.get("OLA")),
            "scraper_status": _clean(r.get("Status")),
            "product_url": _clean(r.get("Product URL")) or _clean(r.get("SKU List Product URL")),
            "raw": {str(k): _json_safe(v) for k, v in r.to_dict().items() if not str(k).lower().startswith("image_")},
        })
    return rows


def _comparison(expected: float, actual: float | None, in_stock: bool | None, scraper_status: str, tolerance_pct: float):
    if scraper_status.lower().startswith("error"):
        return "error", None, None, None
    if actual is None:
        return "missing", None, None, None
    if in_stock is False:
        return "unavailable", None, round(actual - expected, 2), round(((actual / expected) - 1) * 100, 4) if expected else None
    difference = round(actual - expected, 2)
    over_pct = round(((actual / expected) - 1) * 100, 4) if expected else None
    if actual <= expected:
        return "ok", True, difference, over_pct
    if actual <= expected * (1 + tolerance_pct):
        return "half_issue", True, difference, over_pct
    return "issue", False, difference, over_pct


def _start_run(db, account: dict, reference_set_id: str, checkpoint: str, expected: int) -> dict:
    now = _now()
    run_date = now.date().isoformat()
    existing = db.table("cpp_price_runs").select("id,revision").eq("account_id", account["id"]).eq("run_date", run_date).eq("checkpoint", checkpoint).eq("is_current", True).order("revision", desc=True).limit(1).execute().data or []
    revision = (int(existing[0]["revision"]) + 1) if existing else 1
    supersedes = existing[0]["id"] if existing else None
    if existing:
        db.table("cpp_price_runs").update({"is_current": False}).eq("id", supersedes).execute()
    row = db.table("cpp_price_runs").insert({
        "account_id": account["id"],
        "reference_set_id": reference_set_id,
        "run_date": run_date,
        "checkpoint": checkpoint,
        "revision": revision,
        "is_current": True,
        "supersedes_run_id": supersedes,
        "started_at": now.isoformat(),
        "status": "running",
        "rows_expected": expected,
    }).execute().data[0]
    return row


def run_cpp_scrape(accounts: list[str] | None = None, checkpoint: str = "manual") -> dict:
    """Run existing marketplace scrapers against the active CPP reference targets.

    checkpoint: 'am', 'pm', or 'manual'. Re-running the same slot on the same day
    creates a new revision and supersedes the previous current run.
    """
    checkpoint = checkpoint.lower().strip()
    if checkpoint not in {"am", "pm", "manual"}:
        raise ValueError("checkpoint must be am, pm, or manual")

    db = client()
    live_targets = fetch_all(
        "v_cpp_live_targets",
        "target_id,reference_set_id,account_id,account_code,account_name,location_id,location_name,account_sku,product_name,expected_price,tolerance_pct,basepack_name,brand,category,product_url",
    )
    if accounts:
        selected = {x.lower().strip() for x in accounts}
        live_targets = [x for x in live_targets if x["account_code"] in selected]

    by_account: dict[str, list[dict]] = defaultdict(list)
    for target in live_targets:
        if target["account_code"] in SCRAPER_FILES:
            by_account[target["account_code"]].append(target)

    runtime = settings.runtime_dir
    input_dir = runtime / "cpp_sku_lists"
    output_dir = runtime / "cpp_daily_files"
    log_dir = runtime / "cpp_logs"
    for folder in (input_dir, output_dir, log_dir):
        folder.mkdir(parents=True, exist_ok=True)

    input_paths = export_cpp_inputs(live_targets, input_dir)
    results = []

    for code, targets in sorted(by_account.items()):
        if code not in input_paths:
            results.append({"account": code, "status": "skipped", "reason": "no CPP scraper input generated"})
            continue

        account = {"id": targets[0]["account_id"], "code": code, "name": targets[0]["account_name"]}
        reference_set_id = targets[0]["reference_set_id"]
        run = _start_run(db, account, reference_set_id, checkpoint, len(targets))
        stamp = _now().strftime("%Y%m%d_%H%M%S")
        output = output_dir / f"{code}_cpp_{checkpoint}_{stamp}.xlsx"
        log_file = log_dir / f"{code}_cpp_{checkpoint}_{stamp}.log"
        script = settings.legacy_root / "scrapers" / SCRAPER_FILES[code]
        cmd = [sys.executable, str(script), "--sku-list", str(input_paths[code]), "--output", str(output)]

        with log_file.open("w", encoding="utf-8") as log:
            proc = subprocess.run(cmd, cwd=str(settings.legacy_root), stdout=log, stderr=subprocess.STDOUT, text=True)

        if proc.returncode != 0 or not output.exists():
            db.table("cpp_price_runs").update({
                "finished_at": _now().isoformat(),
                "status": "failed",
                "log_file": str(log_file),
                "message": f"scraper exit={proc.returncode}",
                "error_count": len(targets),
            }).eq("id", run["id"]).execute()
            results.append({"account": code, "status": "failed", "log": str(log_file)})
            continue

        observed_rows = _read_scrape_output(output, code)
        location_names = {r["id"]: r["name"] for r in fetch_all("locations", "id,name")}
        obs_map: dict[tuple[str, str], dict] = {}
        for row in observed_rows:
            key = (row["account_sku"], row["branch"] if code == "pandamart" else "")
            # Keep first deterministic row. The scraper already resolves duplicates.
            obs_map.setdefault(key, row)

        checks = []
        counts = Counter()
        for target in targets:
            branch = _normalize_branch(location_names.get(target.get("location_id"), "")) if code == "pandamart" else ""
            obs = obs_map.get((_sku_for_account(code, target["account_sku"]), branch))
            expected = float(target["expected_price"])
            tolerance = float(target.get("tolerance_pct") or 0.01)
            if obs is None:
                status, compliant, diff, over = "missing", None, None, None
                actual = original = None
                in_stock = None
                scraper_status = "Not found in scraper output"
                product_name = target.get("product_name")
                product_url = target.get("product_url")
                evidence = {"source_file": output.name, "checkpoint": checkpoint}
            else:
                actual = obs["actual_price"]
                original = obs["original_price"]
                in_stock = obs["in_stock"]
                scraper_status = obs["scraper_status"]
                status, compliant, diff, over = _comparison(expected, actual, in_stock, scraper_status, tolerance)
                product_name = obs["product_name"] or target.get("product_name")
                product_url = obs["product_url"] or target.get("product_url")
                evidence = {**obs["raw"], "source_file": output.name, "checkpoint": checkpoint}
            counts[status] += 1
            checks.append({
                "run_id": run["id"],
                "target_id": target["target_id"],
                "account_id": account["id"],
                "location_id": target.get("location_id"),
                "checked_at": _now().isoformat(),
                "checked_date": _now().date().isoformat(),
                "account_sku": target["account_sku"],
                "product_name": product_name,
                "expected_price": expected,
                "actual_price": actual,
                "original_price": original,
                "price_difference": diff,
                "over_expected_pct": over,
                "in_stock": in_stock,
                "status": status,
                "compliant": compliant,
                "scraper_status": scraper_status,
                "product_url": product_url,
                "evidence": evidence,
            })

        for batch in chunked(checks, 300):
            db.table("cpp_price_checks").insert(batch).execute()

        error_count = counts["error"]
        final_status = "success" if error_count == 0 else "partial"
        db.table("cpp_price_runs").update({
            "finished_at": _now().isoformat(),
            "status": final_status,
            "rows_observed": len(observed_rows),
            "rows_checked": len(checks),
            "ok_count": counts["ok"],
            "half_issue_count": counts["half_issue"],
            "issue_count": counts["issue"],
            "missing_count": counts["missing"],
            "unavailable_count": counts["unavailable"],
            "error_count": error_count,
            "source_file": output.name,
            "log_file": str(log_file),
            "message": f"CPP {checkpoint}: {len(checks)} checks from {len(observed_rows)} scraper rows",
        }).eq("id", run["id"]).execute()

        results.append({
            "account": code,
            "status": final_status,
            "checkpoint": checkpoint,
            "run_id": run["id"],
            "revision": run["revision"],
            "expected": len(targets),
            "observed": len(observed_rows),
            "checks": len(checks),
            "ok": counts["ok"],
            "half_issue": counts["half_issue"],
            "issue": counts["issue"],
            "missing": counts["missing"],
            "unavailable": counts["unavailable"],
            "errors": error_count,
            "output": str(output),
            "log": str(log_file),
        })

    return {
        "run_date": _now().date().isoformat(),
        "checkpoint": checkpoint,
        "accounts": results,
        "targets": len(live_targets),
    }


def run_cpp_checks() -> dict:
    """Backward-compatible alias. The old latest-observation checker is retired."""
    return run_cpp_scrape(checkpoint="manual")
