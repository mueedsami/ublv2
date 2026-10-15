"""
Pandamart Search Visibility Scraper
=====================================
For each keyword × branch, fetches search results via GraphQL
and records which products appear and at what rank position.

Output: One Excel — Summary sheet + one sheet per keyword.
Each keyword sheet: Rank | Branch1 result | Branch2 result | ...

Install:  pip install requests openpyxl
Run:      python pandamart_search_scraper.py
"""

import argparse
from pathlib import Path

import requests
import time
import random
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── KEYWORDS ──────────────────────────────────────────────────────────────────
# Default keywords are generated from search_keywords.csv.
# You can also pass a keyword file at runtime:
#   python script.py --keywords-file keywords_flat.txt --count 4
DEFAULT_KEYWORDS = [
    'soap',
    'bar soap',
    'barsoap',
    'bar',
    'body wash',
    'bodywash',
    'shower gel',
    'conditioner',
    'condi',
    'detergent',
    'detergent powder',
    'washing powder',
    'detergent liquid',
    'detergent bar',
    'dishwash bar',
    'dishwasher bar',
    'dishwash liquid',
    'dishwash powder',
    'facewash',
    'face wash',
    'face cream',
    'moisturizer',
    'moisturiser',
    'lotion',
    'handwash',
    'hand wash',
    'shampoo',
    'hair mask',
    'hair serum',
    'sunscreen',
    'toilet cleaner',
    'toothpaste',
    'tooth paste',
    'toothpowder',
    'tooth powder',
    'petroleum jelly',
    'vaseline'
]


def load_keywords(keywords_file: str = "keywords_flat.txt", keywords_arg: str = "") -> list[str]:
    if keywords_arg:
        return [k.strip() for k in keywords_arg.split(",") if k.strip()]

    if keywords_file:
        p = Path(keywords_file)
        if p.exists():
            return [line.strip() for line in p.read_text(encoding="utf-8-sig").splitlines() if line.strip()]

    return DEFAULT_KEYWORDS

# ── BRANCHES ──────────────────────────────────────────────────────────────────
BRANCHES = {
    "Uttara":          ("n7ph", "pandamart-uttara"),
    "Gulshan-Niketon": ("w2lx", "pandamart-gulshan-niketon"),
    "Gulshan-Banani":  ("vbpl", "pandamart-gulshan-banani"),
    "Mirpur-Zoo-Road": ("kx59", "pandamart-mirpur-zoo-road"),
    "Mogbazar":        ("v1ts", "pandamart-mogbazar"),
    "Mohammadpur":     ("bbus", "pandamart-mohammadpur"),
    "Rampura":         ("mfg7", "pandamart-rampura"),
    "Dhanmondi":       ("ta7z", "pandamart-dhanmondi"),
    "Wari":            ("zj9s", "pandamart-wari"),
    "Bashundhara":     ("h5rj", "pandamart-bashundhara"),
}

# ── Config ────────────────────────────────────────────────────────────────────
API_URL     = "https://bd.fd-api.com/api/v5/graphql"
RESULTS_PER = 40
DELAY       = (0.4, 0.8)

HEADERS = {
    "User-Agent":                   "Mozilla/5.0 (Linux; Android 6.0; Nexus 5 Build/MRA58N) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Mobile Safari/537.36",
    "Accept":                       "application/json",
    "Content-Type":                 "application/json;charset=UTF-8",
    "accept-language":              "en-US,en;q=0.5",
    "Referer":                      "https://www.foodpanda.com.bd/",
    "Origin":                       "https://www.foodpanda.com.bd",
    "apollographql-client-name":    "web",
    "apollographql-client-version": "GROCERIES-MENU-MICROFRONTEND.26.10.0007",
    "perseus-client-id":            "1771735600556.496036114264219763.2kr7g7xzen",
    "perseus-session-id":           "1772997693938.887648781557212658.qn395k3baj",
    "platform":                     "web",
    "x-pd-language-id":             "1",
    "x-requested-with":             "XMLHttpRequest",
}

GRAPHQL_QUERY = """
    fragment ProductFields on Product {
        attributes(keys: $attributes) {
            key
            value
        }
        isAvailable
        name
        originalPrice
        price
        productID
        stockAmount
        vendorID
    }

    query getSearchProducts(
        $attributes: [String!]
        $brand: String!
        $config: String
        $userCode: String
        $globalEntityId: String
        $languageCode: String!
        $limit: Int!
        $locale: String!
        $offset: Int!
        $query: String!
        $vendors: [VendorInformation!]!
        $verticalTypes: [String!]!
    ) {
        searchProducts(
            input: {
                brand: $brand
                config: $config
                customerID: $userCode
                globalEntityID: $globalEntityId
                languageCode: $languageCode
                limit: $limit
                locale: $locale
                offset: $offset
                query: $query
                vendors: $vendors
                verticalTypes: $verticalTypes
            }
        ) {
            itemCounts { total }
            products {
                items {
                    payload { ...ProductFields }
                }
            }
        }
    }
"""

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
XLSX_FILE = f"pandamart_search_{timestamp}.xlsx"

BRANCH_NAMES = list(BRANCHES.keys())

# ── STYLES ────────────────────────────────────────────────────────────────────
HDR_FILL  = PatternFill("solid", fgColor="1F3864")
HDR_FONT  = Font(bold=True, color="FFFFFF", size=10, name="Arial")
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
NO_FILL   = PatternFill(fill_type=None)
STD_FONT  = Font(size=10, name="Arial")
CTR       = Alignment(horizontal="center", vertical="center", wrap_text=True)
MID       = Alignment(vertical="center", wrap_text=True)
thin      = Side(style="thin", color="CCCCCC")
BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── SEARCH ────────────────────────────────────────────────────────────────────
def search_branch(keyword, vendor_id, vendor_slug):
    variables = {
        "attributes":     ["baseContentValue", "baseUnit", "freshnessGuaranteeInDays",
                           "maximumSalesQuantity", "minPriceLastMonth", "pricePerBaseUnit",
                           "sku", "nutri_grade", "sugar_level"],
        "brand":          "fp",
        "config":         None,
        "globalEntityId": "FP_BD",
        "languageCode":   "en",
        "limit":          RESULTS_PER,
        "locale":         "en_BD",
        "offset":         0,
        "query":          keyword,
        "vendors":        [{"id": vendor_id}],
        "verticalTypes":  ["darkstores"],
    }
    payload = {"query": GRAPHQL_QUERY, "variables": variables}
    try:
        r = requests.post(API_URL, json=payload, headers=HEADERS, timeout=15)
        if r.status_code != 200:
            return [], f"HTTP {r.status_code}"

        data     = r.json()
        errors   = data.get("errors")
        if errors:
            msg = str(errors[0].get("message", errors))
            # "null" error = no results found for this query/branch combo
            if "null" in msg.lower():
                return [], 0
            return [], msg

        items = (data.get("data", {})
                     .get("searchProducts", {})
                     .get("products", {})
                     .get("items", []))
        total = (data.get("data", {})
                     .get("searchProducts", {})
                     .get("itemCounts", {})
                     .get("total", 0))

        results = []
        for rank, item in enumerate(items[:RESULTS_PER], 1):
            p = item.get("payload", {})
            results.append({
                "rank":  rank,
                "name":  p.get("name", "N/A"),
                "price": p.get("price", ""),
                "stock": p.get("stockAmount", 0),
                "ola":   1 if p.get("isAvailable") else 0,
            })
        return results, total

    except Exception as e:
        return [], str(e)

# ── WRITE KEYWORD SHEET ───────────────────────────────────────────────────────
def write_keyword_sheet(wb, keyword, branch_results):
    ws = wb.create_sheet(title=keyword[:31])
    n  = len(BRANCH_NAMES)

    # Title
    ws.merge_cells(f"A1:{get_column_letter(n + 1)}1")
    t = ws.cell(row=1, column=1,
        value=f'Search: "{keyword}"  |  {datetime.now().strftime("%Y-%m-%d")}')
    t.font = Font(bold=True, size=11, name="Arial", color="1F3864")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Headers
    c = ws.cell(row=2, column=1, value="Rank")
    c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = CTR; c.border = BORDER
    ws.column_dimensions["A"].width = 6

    for bi, bname in enumerate(BRANCH_NAMES, 2):
        c = ws.cell(row=2, column=bi, value=bname)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = CTR; c.border = BORDER
        ws.column_dimensions[get_column_letter(bi)].width = 35
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "A3"

    max_results = max((len(v) for v in branch_results.values()), default=0)
    if max_results == 0:
        ws.cell(row=3, column=1, value="No results found").font = Font(italic=True, name="Arial", size=10)
        return

    for ri in range(max_results):
        row_num = ri + 3
        fill    = GREY_FILL if row_num % 2 == 0 else NO_FILL

        c = ws.cell(row=row_num, column=1, value=ri + 1)
        c.font = Font(bold=True, size=10, name="Arial")
        c.fill = fill; c.alignment = CTR; c.border = BORDER

        for bi, bname in enumerate(BRANCH_NAMES, 2):
            results = branch_results.get(bname, [])
            name    = results[ri]["name"] if ri < len(results) else "—"
            c = ws.cell(row=row_num, column=bi, value=name)
            c.font = STD_FONT; c.fill = fill; c.alignment = MID; c.border = BORDER

# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    global RESULTS_PER

    parser = argparse.ArgumentParser(description="Pandamart search visibility scraper")
    parser.add_argument("--keywords-file", "-f", default="keywords_flat.txt")
    parser.add_argument("--keywords", "-k", default="")
    parser.add_argument("--count", "-n", type=int, default=RESULTS_PER)
    args = parser.parse_args()

    keywords = load_keywords(args.keywords_file, args.keywords)
    RESULTS_PER = args.count

    total_calls = len(keywords) * len(BRANCHES)
    print("=" * 60)
    print("  Pandamart Search Visibility Scraper")
    print(f"  Date     : {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print(f"  Keywords : {len(keywords)}")
    print(f"  Branches : {len(BRANCHES)}")
    print(f"  Top N    : {RESULTS_PER}")
    print(f"  Total API calls: {total_calls}")
    print("=" * 60)
    print()

    wb = Workbook()
    wb.remove(wb.active)

    # Summary sheet
    ws_sum = wb.create_sheet("Summary", 0)
    ws_sum.merge_cells(f"A1:{get_column_letter(len(BRANCHES) + 1)}1")
    t = ws_sum.cell(row=1, column=1, value="Pandamart Search Visibility — Result Counts per Branch")
    t.font = Font(bold=True, size=13, name="Arial", color="1F3864")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 28

    c = ws_sum.cell(row=2, column=1, value="Keyword")
    c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = CTR; c.border = BORDER
    ws_sum.column_dimensions["A"].width = 20
    for bi, bname in enumerate(BRANCH_NAMES, 2):
        c = ws_sum.cell(row=2, column=bi, value=bname)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = CTR; c.border = BORDER
        ws_sum.column_dimensions[get_column_letter(bi)].width = 16
    ws_sum.freeze_panes = "A3"

    for ki, keyword in enumerate(keywords, 1):
        print(f"[{ki:>2}/{len(keywords)}] '{keyword}'")
        branch_results = {}

        for bname, (vendor_id, vendor_slug) in BRANCHES.items():
            print(f"    {bname:<18}", end=" → ")
            results, info = search_branch(keyword, vendor_id, vendor_slug)
            if isinstance(info, str):
                print(f"✗ {info}")
                results = []
            else:
                print(f"✓ {len(results)} results")
            branch_results[bname] = results
            time.sleep(random.uniform(*DELAY))

        write_keyword_sheet(wb, keyword, branch_results)

        # Summary row
        sr   = ki + 2
        fill = GREY_FILL if sr % 2 == 0 else NO_FILL
        c = ws_sum.cell(row=sr, column=1, value=keyword)
        c.font = Font(bold=True, size=10, name="Arial")
        c.fill = fill; c.alignment = MID; c.border = BORDER
        for bi, bname in enumerate(BRANCH_NAMES, 2):
            count = len(branch_results.get(bname, []))
            c = ws_sum.cell(row=sr, column=bi, value=count)
            c.font = STD_FONT; c.fill = fill; c.alignment = CTR; c.border = BORDER

        print()

    wb.save(XLSX_FILE)
    print(f"{'='*60}")
    print(f"  Done! Saved to: {XLSX_FILE}")
    print(f"{'='*60}")

if __name__ == "__main__":
    main()